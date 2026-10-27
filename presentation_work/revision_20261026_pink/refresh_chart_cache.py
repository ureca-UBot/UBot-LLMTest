"""Refresh chart caches from their existing exact worksheet ranges.
PowerPoint serializes 9.13 as 9.1300000000000008; retain the embedded data.
"""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from io import BytesIO
from lxml import etree as E
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
import posixpath as pp,json,re
D=Path(__file__).resolve().parent
N={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
with ZipFile(D/'integrated_com_candidate_v2.pptx') as z:parts={n:z.read(n) for n in z.namelist()}
changes=[]
for name in list(parts):
 if not re.fullmatch(r'ppt/charts/chart\d+\.xml',name):continue
 root=E.fromstring(parts[name]);external=root.find('c:externalData',N)
 if external is None:continue
 rels=E.fromstring(parts[pp.join(pp.dirname(name),'_rels',pp.basename(name)+'.rels')])
 target=next(r.get('Target') for r in rels if r.get('Id')==external.get('{'+N['r']+'}id'))
 bookpart=pp.normpath(pp.join(pp.dirname(name),target))
 wb=load_workbook(BytesIO(parts[bookpart]),data_only=True)
 for ref in root.xpath('//c:numRef|//c:strRef',namespaces=N):
  formula=ref.find('c:f',N)
  if formula is None:continue
  sheet,ran=formula.text.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
  a,b,c,d=range_boundaries(ran);values=[wb[sheet].cell(row,col).value for row in range(b,d+1) for col in range(a,c+1)]
  cache=ref.find('c:numCache',N) if E.QName(ref).localname=='numRef' else ref.find('c:strCache',N)
  if cache is None:continue
  for point in cache.findall('c:pt',N):
   val=values[int(point.get('idx'))];text='' if val is None else str(val)
   v=point.find('c:v',N)
   if v is not None and v.text!=text:changes.append(dict(chart=name,range=formula.text,index=point.get('idx'),before=v.text,after=text));v.text=text
 parts[name]=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(D/'integrated_fixed_candidate_v2.pptx','w',ZIP_DEFLATED) as z:
 for name,data in parts.items():z.writestr(name,data)
(D/'chart_cache_refresh.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(changes,ensure_ascii=False))
